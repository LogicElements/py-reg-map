"""Import a legacy Excel/VBA register map workbook (e.g. Vms1511.xlsm) into YAML."""

import os
import re
import warnings
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException

from regmap.importers.yaml_emit import HexInt, emit_map
from regmap.model import MapError, load_map_text
from regmap.resolve import resolve

REGISTER_COLUMNS = {
    "name_source": "name source",
    "block": "block code",
    "type": "data type",
    "access": "access code",
    "address": "address",
    "size": "max length",
    "id": "id [hex]",
    "name": "name",
    "label": "label",
    "default": "factory value",
    "min": "minimal value",
    "max": "maximal value",
    "unit": "units",
    "modbus": "modbus special",
    "invalid": "invalid",
    "description": "description",
}
BLOCK_COLUMNS = {"abbrev": "abbrev.", "code": "code", "name": "name", "description": "description"}
DESTINATIONS = {
    "regmap destination": ("reg_map",),
    "modbus destination": ("modbus",),
    "tests destination": ("lebin_json", "modbus_json", "python"),
}
INPUT_ACCESS = ("RO", "ROF")
_HEX = re.compile(r"^0[xX][0-9A-Fa-f]+$")
_GETPATH = re.compile(r'^=\s*GetPath\(\)\s*&\s*"(.*)"\s*$', re.IGNORECASE)
_IDENT = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")


class ImportFailed(Exception):
    def __init__(self, errors: list[str]) -> None:
        super().__init__("\n".join(errors))
        self.errors = errors


@dataclass(frozen=True)
class ImportResult:
    yaml_text: str
    warnings: list[str]
    register_count: int


def _norm(value: Any) -> str:
    return " ".join(str(value).split()).lower() if value is not None else ""


def _text(value: Any) -> str | None:
    """Cell as text: lines right-trimmed, surrounding empty lines removed; None when empty."""
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    lines = str(value).replace("\r\n", "\n").replace("\r", "\n").split("\n")
    text = "\n".join(line.rstrip() for line in lines).strip("\n")
    return text if text.strip() else None


def _number(value: Any) -> int | float | None:
    """Cell as number (hex text kept as HexInt); None when empty or not a number."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value) if value.is_integer() else value
    text = str(value).strip()
    if _HEX.match(text):
        return HexInt(text)
    for convert in (int, lambda t: float(t.replace(",", "."))):
        try:
            return convert(text)
        except ValueError:
            pass
    return None


def _open(path: Path, data_only: bool):
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")  # openpyxl warns about unsupported Excel extensions
        return openpyxl.load_workbook(path, read_only=True, data_only=data_only)


def _rows(sheet) -> list[tuple[Any, ...]]:
    return [tuple(row) for row in sheet.iter_rows(values_only=True)]


def _columns(header: tuple[Any, ...], wanted: dict[str, str], where: str) -> dict[str, int]:
    found = {_norm(v): i for i, v in reversed(list(enumerate(header))) if v is not None}
    missing = [text for text in wanted.values() if text not in found]
    if missing:
        raise ImportFailed([f"{where}: missing column(s): {', '.join(missing)}"])
    return {key: found[text] for key, text in wanted.items()}


def _cell(row: tuple[Any, ...], index: int) -> Any:
    return row[index] if index < len(row) else None


def _label_value(rows: list[tuple[Any, ...]], label: str) -> Any:
    """Value of the first non-empty cell right of a cell whose text is ``label``."""
    for row in rows:
        for i, value in enumerate(row):
            if _norm(value) == label:
                return next((v for v in row[i + 1 :] if v not in (None, "")), None)
    return None


def _read_blocks(sheet) -> dict[str, dict[str, Any]]:
    rows = _rows(sheet)
    start = next((i for i, r in enumerate(rows) if any(_norm(v) == "block id" for v in r)), None)
    if start is None:
        raise ImportFailed(["Common sheet: 'Block ID' table not found"])
    head = next(
        (i for i in range(start, len(rows)) if any(_norm(v) == "abbrev." for v in rows[i])), None
    )
    if head is None:
        raise ImportFailed(["Common sheet: header row with 'Abbrev.' not found"])
    cols = _columns(rows[head], BLOCK_COLUMNS, "Common sheet")
    blocks = {}
    for row in rows[head + 1 :]:
        code = _number(_cell(row, cols["code"]))
        if code is None:
            break
        abbrev = _text(_cell(row, cols["abbrev"]))
        if abbrev:
            blocks[abbrev] = {
                "code": int(code),
                "name": _text(_cell(row, cols["name"])),
                "description": _text(_cell(row, cols["description"])),
            }
    return blocks


def _vba_word_count(reg_type: str, size: int, f32: bool) -> int:
    """Modbus words as the VBA macro counted them (banker's rounding of size / 2)."""
    if reg_type in ("INT", "STRING", "BIN"):
        return max(round(size / 2), 1)
    if reg_type == "FLOAT":
        return 2 if f32 else 1
    return 1


def _word_count(reg_type: str, size: int, f32: bool) -> int:
    if reg_type == "ENUM":
        return 1
    if reg_type == "FLOAT":
        return 2 if f32 else 1
    return (size + 1) // 2


def _destination(formula: Any, workbook_dir: Path, yaml_dir: Path) -> str | None:
    match = _GETPATH.match(str(formula).strip())
    if not match:
        return None
    relative = match.group(1).replace("\\", "/").strip("/")
    target = os.path.normpath(workbook_dir / relative)
    try:
        return Path(os.path.relpath(target, yaml_dir)).as_posix()
    except ValueError:  # different drive on Windows
        return None


class _Importer:
    def __init__(self, workbook: Path, target: Path) -> None:
        self.workbook = workbook
        self.target = target
        self.warnings: list[str] = []
        self.errors: list[str] = []
        self.excel_ids: dict[str, int] = {}
        self.vba_modbus: dict[str, tuple[str, int, int]] = {}  # full name -> (space, start, count)
        self.vba_counter = {"INPUT": 0, "HOLD": 0}

    def run(self) -> ImportResult:
        try:
            values = _open(self.workbook, data_only=True)
            formulas = _open(self.workbook, data_only=False)
        except (zipfile.BadZipFile, InvalidFileException, KeyError) as exc:
            raise ImportFailed([f"cannot read the workbook: {exc}"]) from None
        for name in ("Registers", "Common"):
            if name not in values.sheetnames:
                raise ImportFailed([f"sheet '{name}' not found"])
        rows = _rows(values["Registers"])
        header = next(
            (i for i, r in enumerate(rows) if any(_norm(v) == "block code" for v in r)), None
        )
        if header is None:
            raise ImportFailed(["Registers sheet: header row with 'Block Code' not found"])
        cols = _columns(rows[header], REGISTER_COLUMNS, "Registers sheet")
        known_blocks = _read_blocks(values["Common"])
        registers = self._read_registers(rows, header, cols, known_blocks)
        if self.errors:
            raise ImportFailed(self.errors)
        doc = {
            "device": self._device(rows[:header]),
            "generator": self._generator(_rows(formulas["Registers"])[:header]),
            "blocks": self._blocks(registers, known_blocks),
        }
        text = emit_map(doc)
        self._verify(text)
        return ImportResult(text, self.warnings, sum(len(r) for r in registers.values()))

    # -- sheet rows -> plain register dicts

    def _read_registers(self, rows, header, cols, known_blocks) -> dict[str, list[dict[str, Any]]]:
        registers: dict[str, list[dict[str, Any]]] = {}
        offsets: dict[str, int] = {}
        current: dict[str, Any] | None = None
        skipping = False  # items of a skipped (Invalid) register are dropped silently
        for number, row in enumerate(rows[header + 1 :], start=header + 2):

            def get(key: str, row=row) -> Any:
                return _cell(row, cols[key])

            special = _text(get("modbus")) or ""
            self._vba_special(special)
            reg_type = _text(get("type"))
            name_source = _text(get("name_source"))
            if reg_type:  # a register row, or a separator such as "stop"
                current, skipping = None, False
                if not name_source:
                    continue
                if _text(get("invalid")):
                    self.warnings.append(f"row {number}: {name_source} is marked Invalid, skipped")
                    skipping = True
                    continue
                block = _text(get("block"))
                if block not in known_blocks:
                    self.errors.append(f"row {number}: block {block!r} is not in the Common sheet")
                    continue
                current = self._register(
                    number, get, block, reg_type, name_source, special, offsets
                )
                if current is not None:
                    registers.setdefault(block, []).append(current)
            elif _text(get("name")):
                if current is not None:
                    self._item(number, get, current)
                elif not skipping:
                    self.warnings.append(
                        f"row {number}: item {_text(get('name'))} has no register above it, ignored"
                    )
        for regs in registers.values():
            for reg in regs:
                self._finish_items(reg)
        return registers

    def _vba_special(self, special: str) -> None:
        parts = special.split()
        if len(parts) == 2 and parts[0] in ("IA", "HA", "IR", "HR") and parts[1].isdigit():
            space = "INPUT" if parts[0][0] == "I" else "HOLD"
            n = int(parts[1])
            self.vba_counter[space] = n if parts[0][1] == "A" else self.vba_counter[space] + n

    def _register(self, number, get, block, reg_type, name, special, offsets):
        full = f"{block}_{name}"
        size = _number(get("size"))
        address = _number(get("address"))
        if size is None or address is None:
            self.errors.append(f"row {number}: {full} needs Address and Max length")
            return None
        size, address = int(size), int(address)
        access = _text(get("access")) or ""
        reg: dict[str, Any] = {"name": name, "type": reg_type, "access": access, "size": size}
        offset = offsets.get(block, 0)
        if address != offset:
            reg["address"] = address
        offsets[block] = max(offset, address + size)
        label = _text(get("label"))
        if label and "\n" in label:
            self.warnings.append(f"row {number}: multi-line label of {full} joined into one line")
            label = " ".join(label.split("\n"))
        if label:
            reg["label"] = label
        raw_default = get("default")
        if reg_type == "STRING":
            if _text(raw_default):
                reg["default"] = _text(raw_default)
        elif raw_default not in (None, ""):
            value = _number(raw_default)
            if value is None:
                self.warnings.append(f"row {number}: default of {full} is not a number, dropped")
            else:
                reg["default"] = value
        for key in ("min", "max"):
            value = _number(get(key))
            if value is not None:
                reg[key] = value
        unit = _text(get("unit"))
        if unit:
            reg["unit"] = unit
        f32 = special == "F32"
        if special.lower() == "x":
            reg["modbus"] = False
        elif f32 and reg_type != "FLOAT":
            self.warnings.append(f"row {number}: F32 on non-FLOAT register {full} ignored")
            f32 = False
        elif f32:
            reg["modbus"] = {"format": "F32"}
        elif special and not re.match(r"^(IA|HA|IR|HR) \d+$", special):
            self.warnings.append(f"row {number}: unknown Modbus special {special!r} ignored")
        description = _text(get("description"))
        if description:
            reg["description"] = description
        reg["_items"] = []
        reg["_row"] = number
        excel_id = _text(get("id"))
        if excel_id and _HEX.match(excel_id):
            self.excel_ids[full] = int(excel_id, 16)
        if reg.get("modbus") is not False:
            space = "INPUT" if access in INPUT_ACCESS else "HOLD"
            count = _vba_word_count(reg_type, size, f32)
            self.vba_modbus[full] = (space, self.vba_counter[space], count)
            self.vba_counter[space] += count
        return reg

    def _item(self, number, get, reg) -> None:
        if reg["type"] not in ("ENUM", "BIN"):
            self.warnings.append(
                f"row {number}: item {_text(get('name'))} under {reg['type']} register ignored"
            )
            return
        reg["_items"].append(
            {
                "name": _text(get("name")),
                "label": _text(get("label")),
                "description": _text(get("description")),
                "override": _number(get("default")),
            }
        )

    def _finish_items(self, reg: dict[str, Any]) -> None:
        items = reg.pop("_items")
        row = reg.pop("_row")
        running = 0
        converted = []
        for item in items:
            number = running if item["override"] is None else int(item["override"])
            entry = {"name": item["name"]}
            if reg["type"] == "BIN":
                entry["bit"] = number
            elif number != running:
                entry["value"] = number
            entry["label"] = item["label"]
            entry["description"] = item["description"]
            converted.append((entry, number))
            running = number + 1
        if reg["type"] == "BIN" and converted:
            reg["bits"] = [e for e, _ in converted]
        if reg["type"] != "ENUM":
            return
        reg["values"] = [e for e, _ in converted]
        numbers = [n for _, n in converted]
        full = f"row {row}: {reg['name']}"
        if "default" in reg:
            by_number = {n: e["name"] for e, n in converted}
            if reg["default"] in by_number:
                reg["default"] = by_number[reg["default"]]
            else:
                self.warnings.append(f"{full}: default {reg['default']} is not a value, dropped")
                del reg["default"]
        sheet_range = (reg.pop("min", None), reg.pop("max", None))
        if numbers and sheet_range != (None, None) and sheet_range != (min(numbers), max(numbers)):
            self.warnings.append(
                f"{full}: sheet min/max {sheet_range} differ from the values, the values win"
            )

    # -- document parts

    def _device(self, top_rows) -> dict[str, Any]:
        name = self.workbook.stem
        if not _IDENT.match(name):
            raise ImportFailed([f"workbook name {name!r} is not a valid device name"])
        device: dict[str, Any] = {"name": name}
        prefix = _text(_label_value(top_rows, "config c prefix"))
        storage = _text(_label_value(top_rows, "config storage"))
        if prefix and prefix != "CONF_":
            device["c_prefix"] = prefix
        if storage and storage != "CONF_REG":
            device["c_storage"] = storage
        return device

    def _generator(self, top_rows) -> dict[str, Any]:
        outputs: dict[str, str] = {}
        workbook_dir = self.workbook.resolve().parent
        yaml_dir = self.target.resolve().parent
        for label, keys in DESTINATIONS.items():
            formula = _label_value(top_rows, label)
            if formula is None:
                continue
            path = _destination(formula, workbook_dir, yaml_dir)
            if path is None:
                self.warnings.append(f"'{label}' is not a GetPath() formula, not imported")
                continue
            for key in keys:
                outputs[key] = path
        return {"outputs": outputs} if outputs else {}

    def _blocks(self, registers, known_blocks) -> dict[str, Any]:
        self._modbus_overrides(registers, known_blocks)
        blocks = {}
        for abbrev in sorted(registers, key=lambda a: known_blocks[a]["code"]):
            info = known_blocks[abbrev]
            block = {"code": info["code"], "name": info["name"], "description": info["description"]}
            blocks[abbrev] = {
                **{k: v for k, v in block.items() if v is not None},
                "registers": registers[abbrev],
            }
        return blocks

    def _modbus_overrides(self, registers, known_blocks) -> None:
        """Pin Modbus start addresses wherever automatic allocation would differ from VBA."""
        counter = {"INPUT": 0, "HOLD": 0}
        for abbrev in sorted(registers, key=lambda a: known_blocks[a]["code"]):
            for reg in registers[abbrev]:
                vba = self.vba_modbus.get(f"{abbrev}_{reg['name']}")
                if vba is None:
                    continue
                space, start, _ = vba
                if counter[space] != start:
                    options = reg["modbus"] if isinstance(reg.get("modbus"), dict) else {}
                    reg["modbus"] = {"address": start, **options}
                    counter[space] = start
                f32 = isinstance(reg.get("modbus"), dict) and reg["modbus"].get("format") == "F32"
                counter[space] += _word_count(reg["type"], reg["size"], f32)

    def _verify(self, text: str) -> None:
        try:
            resolved = resolve(load_map_text(text, str(self.target)))
        except MapError as exc:
            raise ImportFailed(["the converted map is not valid:"] + exc.errors) from None
        for reg in resolved.registers:
            expected = self.excel_ids.get(reg.full_name)
            if expected is not None and expected != reg.id:
                self.errors.append(
                    f"{reg.full_name}: computed ID 0x{reg.id:08X}"
                    f" differs from the sheet ID 0x{expected:08X}"
                )
            vba = self.vba_modbus.get(reg.full_name)
            if (
                vba
                and reg.modbus
                and (reg.modbus.addresses[0], len(reg.modbus.addresses)) != vba[1:]
            ):
                self.warnings.append(
                    f"{reg.full_name}: Modbus words {list(reg.modbus.addresses)} differ from the"
                    f" VBA allocation (start {vba[1]}, {vba[2]} words)"
                )
        if self.errors:
            raise ImportFailed(self.errors)


def import_workbook(workbook: Path, target: Path) -> ImportResult:
    """Convert ``workbook``; ``target`` is where the YAML will be written (for relative paths)."""
    return _Importer(workbook, target).run()
